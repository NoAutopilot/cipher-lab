#!/usr/bin/env python3
"""holder_export.py: turn a list of read items into a spreadsheet a holding library can load against its own records.

Reads a list TSV like outreach/huntington-decipherments-list-2026-10-08.tsv (comment lines start with '#'; columns
'document', 'depth (rule 4a)', 'depth_pct', 'unresolved', a 'weak ...' column, 'documents' and 'folder'), joins each
row to its status.json result (by document_id) and to the target folder named in 'folder', and writes ONE ROW PER
ITEM (telegram, letter, enclosure) keyed by the holder's own page or object pointer, so the holder can match it to
its catalogue records (a CONTENTdm number, for example). Output: STEM.csv and STEM.tsv always (UTF-8; no tab or
line break inside a cell), STEM.xlsx when openpyxl is importable (header frozen, text wrapped, a README sheet that
explains every column); without openpyxl it says so and writes the two text files only.

Columns, in this order (COLUMNS): holder call number; holder page (as the holder titles it); holder pointer;
holder URL; date; from; to; place; ledger/page text as we transcribed it; reading; code words and meanings;
uncertain or unread words; how much is read; adds less? (and why); prior print checked; partly in print?;
our class; repository link; suggested credit line.

Target layouts read (auto-detected, no per-target configuration):
  md-blocks   reading*.md files with a '<!-- decode.py: derived block starts -->' block of '**ID | ...**' paragraphs
              and the matching ciphertext*.txt '### ID | ...' blocks and key*.md (same suffix: reading-no2.md <->
              ciphertext-no2.txt <-> key-no2.md); the code-word list is rebuilt per token with the folder's own
              decode.py (load_key, load_ciphertext, entry_text, lookup) -- never a private copy of its rules -- and its
              grade counts are checked against the reading's 'Code-word tokens:' line (a mismatch is warned).
              Holder page titles come from the folder's cached CONTENTdm records sources/**/p<pointer>.json.
  decode-key  decode.json jobs of tools/decode_key.py (ciphertext TSV, reading TSV, token TSV with grades); an item is a
              line-id prefix (BLA186 in BLA186_p1_L01), matched against the document text with spaces removed;
              page pointers from images/manifest.json (item, page, pointer) when present.

Splitting: a row whose status.json 'documents' name two items the reading files keep apart (E4 and E5) becomes two
rows. A row whose 'documents' count is larger than the items found (two telegrams in one ledger entry) stays one row
and its reading cell says 'N telegrams in this entry'.

--meta FILE: TSV with an 'item' column and any of holder_call, holder_page, pointer, url, date, from, to, place,
prior_print, in_print, note (a 'source' column is allowed and ignored); a non-empty cell replaces the computed value
('-' blanks it), 'note' is put in front of the reading. Use it for what the folder does not state in a parseable form (the Blathwayt object pointers, a sender).

--select-*: re-run a selection rule on status.json and report drift against the list (rows that no longer pass, and
results that pass but are not listed). Report only: the list is never edited, and the export follows the list.

Wording rules (CLAUDE.md rules 4, 4a, 10): grades spelled out (H from the cipher book, C from a printed or period
copy, S cryptanalytic, I inferred, M uncertain, U unread); depth by the rule-4a outward words; class by its N-class
meaning; the prior-print cell says 'not found in ...' and never new/first/unpublished. Personal data (rule 9): the
tool writes only the list, status.json and the folder's files; nothing from the environment.

Must catch (tests/test_holder_export.py): column order; the E4/E5-style split; the 'N telegrams in this entry' note;
the fallback without openpyxl; no '[SIGN-OFF]' and no e-mail address in any output. Must NOT block: a list row with no
status.json result or no readable folder (written with the list's own fields, warned on stderr).

Usage:
  python3 tools/holder_export.py LIST.tsv --out DIR/STEM [--meta META.tsv] [--status status.json]
      [--url-template 'https://host/digital/collection/{alias}/id/{pointer}'] [--alias mssEC=p16003coll11 ...]
      [--holder 'The Huntington Library'] [--credit TEXT] [--no-xlsx]
      [--select-class N3,N4 --select-min-audits 2 --select-scope completed-reading,recovered-passages
       --select-folder eckert-1864,huntington-blathwayt-madrid-1728]
Exit 0 on success (drift is reported, not fatal); 2 on a usage error.
"""
import argparse
import csv
import datetime
import importlib.util
import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
REPO_URL = "https://github.com/NoAutopilot/cipher-lab"

COLUMNS = [
    "holder call number",
    "holder page (as the holder titles it)",
    "holder pointer (CONTENTdm number)",
    "holder URL",
    "date",
    "from",
    "to",
    "place",
    "ledger/page text as we transcribed it",
    "reading",
    "code words and meanings",
    "uncertain or unread words",
    "how much is read",
    "adds less? (and why)",
    "prior print checked",
    "partly in print?",
    "our class",
    "repository link",
    "suggested credit line",
]
META_KEYS = {"holder_call": 0, "holder_page": 1, "pointer": 2, "url": 3, "date": 4, "from": 5, "to": 6, "place": 7,
             "prior_print": 14, "in_print": 15}

GRADE_LABEL = {"H": "from the cipher book", "C": "from a printed or period copy", "S": "cryptanalytic",
               "I": "inferred", "M": "uncertain", "U": "unread"}
GRADE_ORDER = "HCSIMU"
CLASS_MEANING = {
    "N0": "the plaintext and a decipherment of this very item were already known",
    "N1": "the plaintext is already published; ours is an independent re-decipherment",
    "N2": "the plaintext is known elsewhere, but no prior mapping of this cipher text to it was found",
    "N3": "no prior plaintext or decipherment located after a logged search",
    "N4": "no prior decipherment located, with the principal editions, catalogues and project pages searched "
          "(internal or unpublished work not excluded)",
    "N5": "confirmed by the holding archive or a specialist",
}
NUMBER_WORDS = {"one": 1, "two": 2, "three": 3, "four": 4, "five": 5, "six": 6}

START = "<!-- decode.py: derived block starts -->"
END = "<!-- decode.py: derived block ends -->"
MONTH = r"(?:Jan|Feb|Mar|Apr|May|June?|July?|Aug|Sept?|Oct|Nov|Dec)[a-z]*\.?"
DATE_RE = re.compile(r"\b(\d{1,2}(?:-\d{1,2})? " + MONTH + r" \d{4})")
TIME_RE = re.compile(r"^\s*(\d{1,2}(?:[.:]\d{2})?\s?(?:AM|PM|A\.\s?M\.|P\.\s?M\.|M\b|noon)|midnight|noon)", re.I)
LEDGER_MONTH = re.compile(r"\b(?:Jan|Jany|Feb|Feby|Mar|Mch|Apr|Apl|April|May|June|July|Jul|Aug|Sept|Sep|Oct|Octo|"
                          r"Nov|Novr|Dec|Decr)\b\.?", re.I)
WASH_RE = re.compile(r"\bWash(?:ington|'n|n|\.)?(?=[\s.,]|$)")
NOTE_PREFIXES = ("plain:", "variant:", "split:", "plain-at:", "gloss:", "join:", "note:", "#", "<!--")
NEG_RE = re.compile(r"\b(no prior|not located|not found|neither|unprinted|not printed|nothing|unread|not in the)\b", re.I)
PRINT_RE = re.compile(r"\bprints\b|\bin print\b|\bprinted\b(?!\s+(?:correspondence|papers|editions?|volumes?|works|"
                      r"sources|text|accounts?|prints?)\b)", re.I)
JARGON_RE = re.compile(r"\bAUD|\baudit|\bN[0-5]\b|code clause|\bheld\b|\bweak|\b(?:LS|FM|V1|G3|PROP|FIX)-?[A-Z0-9]|"
                       r"be-api|snippet|decoded|verifier|grade [A-Z]\b", re.I)
CITE_RE = re.compile(r"\bp\.\s?\d|\bpp\.\s?\d|\b1[5-9]\d\d\b|\bORN?\b|Official Records")
ABBR = {"ser", "vol", "vols", "pt", "p", "pp", "no", "nos", "st", "gen", "col", "capt", "lt", "maj", "mr", "dr", "u", "s",
        "jr", "ed", "eds", "cf", "corr", "rec", "i", "ii", "iii", "iv", "v", "a", "b", "c", "d", "e", "f", "g", "h", "j",
        "k", "l", "m", "n", "o", "q", "r", "t", "w", "x", "y", "z", "brig", "lieut", "govr", "esq", "misc", "doc",
        "sess", "cong", "hd", "qrs", "ft", "mss"}


def warn(msg):
    sys.stderr.write("holder_export: " + msg + "\n")


def clean(s):
    """One line per cell: no tab, no line break (a tab-delimited import splits on both)."""
    return re.sub(r"\s+", " ", str(s or "")).strip()


# ---------------------------------------------------------------- list, status, meta

def read_list(path):
    lines = [l for l in Path(path).read_text(encoding="utf-8").splitlines() if l.strip() and not l.startswith("#")]
    if not lines:
        sys.exit("holder_export: empty list " + str(path))
    rows = list(csv.reader(lines, delimiter="\t", quoting=csv.QUOTE_NONE))
    head = rows[0]

    def col(prefix):
        for i, h in enumerate(head):
            if h.strip().lower().startswith(prefix):
                return i
        return None
    idx = {k: col(k) for k in ("document", "title", "class", "depth (", "depth_pct", "completeness", "unresolved",
                               "weak", "key", "documents", "folder")}
    if idx["document"] is None or idx["folder"] is None:
        sys.exit("holder_export: list needs 'document' and 'folder' columns")
    out = []
    for r in rows[1:]:
        r = r + [""] * (len(head) - len(r))
        out.append({k: (r[i].strip() if i is not None else "") for k, i in idx.items()})
    return out


def read_meta(path):
    if not path:
        return {}
    rows = [l for l in Path(path).read_text(encoding="utf-8").splitlines() if l.strip() and not l.startswith("#")]
    rd = list(csv.DictReader(rows, delimiter="\t", quoting=csv.QUOTE_NONE))
    return {r["item"].strip(): {k: (v or "").strip() for k, v in r.items() if k} for r in rd if r.get("item")}


def folder_path(folder, repo_root):
    if "/tree/" in folder:
        rel = folder.split("/tree/", 1)[1].split("/", 1)[1] if folder.split("/tree/", 1)[1].count("/") else ""
        return Path(repo_root) / rel
    return Path(repo_root) / folder


# ---------------------------------------------------------------- md-blocks layout

def _load_module(path):
    spec = importlib.util.spec_from_file_location("holder_export_decoder_" + str(abs(hash(str(path)))), path)
    mod = importlib.util.module_from_spec(spec)
    sys.path.insert(0, str(path.parent))
    try:
        spec.loader.exec_module(mod)
    finally:
        sys.path.pop(0)
    return mod


class MdBlocks:
    """reading*.md derived blocks + ciphertext*.txt blocks + key*.md, read with the folder's own decode.py."""

    def __init__(self, folder):
        self.folder = folder
        self.decoder = _load_module(folder / "decode.py") if (folder / "decode.py").exists() else None
        self.items = {}      # id -> dict(header, lines, reading, counts, key)
        for rpath in sorted(folder.glob("reading*.md")):
            text = rpath.read_text(encoding="utf-8")
            if START not in text or END not in text:
                continue
            suffix = rpath.stem[len("reading"):]
            cpath, kpath = folder / f"ciphertext{suffix}.txt", folder / f"key{suffix}.md"
            if not cpath.exists():
                continue
            key = self.decoder.load_key(kpath) if (self.decoder and kpath.exists()) else None
            if self.decoder:  # the folder's own loader: same block and comment rules as its readings
                blocks = {h.split("|")[0].strip(): (h, ls) for h, ls in self.decoder.load_ciphertext(cpath)}
            else:
                blocks = self._blocks(cpath)
            for iid, para, counts in self._readings(text.split(START, 1)[1].split(END, 1)[0]):
                if iid in blocks:
                    header, lines = blocks[iid]
                    self.items[iid] = {"header": header, "lines": lines, "reading": para, "counts": counts,
                                       "key": key, "reading_file": rpath.name, "cipher_file": cpath.name}

    @staticmethod
    def _blocks(path):
        out, cur = {}, None
        for raw in path.read_text(encoding="utf-8").splitlines():
            if raw.startswith("### "):
                header = raw[4:].strip()
                cur = header.split("|")[0].strip()
                out[cur] = (header, [])
            elif cur is not None and not raw.startswith(("#", "<!--", "note:")):
                out[cur][1].append(raw.rstrip())
        return out

    @staticmethod
    def _readings(block):
        parts = re.split(r"^\*\*(.+?)\*\*\s*$", block, flags=re.M)
        for i in range(1, len(parts) - 1, 2):
            iid = parts[i].split("|")[0].strip()
            body = parts[i + 1]
            m = re.search(r"^Code-word tokens:\s*(.*?)\.?\s*$", body, flags=re.M)
            counts = {}
            if m:
                for g, n in re.findall(r"\b([A-Z]) (\d+)", m.group(1)):
                    counts[g] = int(n)
                body = body[:m.start()]
            para = " ".join(l.strip() for l in body.strip().splitlines() if l.strip())
            yield iid, para, counts

    def match(self, text):
        found = []
        for iid in self.items:
            m = re.search(r"(?<![\w-])" + re.escape(iid) + r"(?![\w-])", text)
            if m:
                found.append((m.start(), iid))
        return [i for _, i in sorted(found)]

    def tokens(self, iid):
        """[(code word as written, meaning, grade)] for every code-word token, in order (decode.py's own lookup)."""
        it = self.items[iid]
        if not (self.decoder and it["key"]):
            return []
        d, key = self.decoder, it["key"]
        out = []
        for w in d.entry_text(it["lines"]).split(" "):
            if w == "|" or not w:
                continue
            if "~" in w:
                surface, target, ovr = w.split("~")
                if target.startswith("!"):
                    out.append((surface.strip(" .,;:'\"()\\"), target[1:].replace("_", " "), ovr or "M"))
                    continue
                stem, flag, row = d.lookup(target.strip(" .,;:'\"()"), key, True)
                if row is not None and ovr:
                    row = (row[0], ovr, row[2])
                stem = surface.strip(" .,;:'\"()\\")
            else:
                stem, flag, row = d.lookup(w.strip(" .,;:'\"()"), key, False)
            if row is None or row[2] in ("blind", "line"):
                continue
            out.append((stem.replace("[?]", "").strip(" \\"), row[0], row[1]))
        return out

    def ledger_text(self, iid):
        lines = [l for l in self.items[iid]["lines"] if l.strip() and not l.startswith(NOTE_PREFIXES)]
        return " / ".join(l.strip() for l in lines)

    def word_count(self, iid):
        it = self.items[iid]
        if not self.decoder:
            return 0
        return len([w for w in self.decoder.entry_text(it["lines"]).split(" ") if w and w != "|"])


# ---------------------------------------------------------------- decode-key layout

class DecodeKey:
    """tools/decode_key.py jobs: ciphertext TSV (line, pos, group), reading TSV (line, text), tokens TSV (grades)."""

    def __init__(self, folder):
        self.folder = folder
        self.items = {}
        cfg = json.loads((folder / "decode.json").read_text(encoding="utf-8"))
        for job in cfg.get("jobs", []):
            tok = folder / job.get("tokens", "")
            rdg = folder / job.get("reading", "")
            cip = folder / job.get("ciphertext", "")
            if not (job.get("tokens") and tok.exists() and rdg.exists()):
                continue
            for r in csv.DictReader(tok.read_text(encoding="utf-8").splitlines(), delimiter="\t"):
                item = r["line"].split("_")[0]
                it = self.items.setdefault(item, {"tokens": [], "reading": [], "cipher": {}, "pages": []})
                it["tokens"].append((r.get("group", ""), r.get("value", ""), r.get("grade", "")))
                page = r["line"].split("_")[1] if r["line"].count("_") >= 2 else ""
                if page and page not in it["pages"]:
                    it["pages"].append(page)
            for l in rdg.read_text(encoding="utf-8").splitlines():
                if l.startswith("#") or "\t" not in l:
                    continue
                lid, txt = l.split("\t", 1)
                if lid.split("_")[0] in self.items:
                    self.items[lid.split("_")[0]]["reading"].append(txt.strip())
            if cip.exists():
                for r in csv.DictReader([l for l in cip.read_text(encoding="utf-8").splitlines()
                                         if not l.startswith("#")], delimiter="\t"):
                    item = r["line"].split("_")[0]
                    if item in self.items:
                        self.items[item]["cipher"].setdefault(r["line"], []).append(r.get("group", ""))
        self.manifest = {}
        man = folder / "images" / "manifest.json"
        if man.exists():
            try:
                for m in json.loads(man.read_text(encoding="utf-8")):
                    if isinstance(m, dict) and m.get("item") and m.get("pointer"):
                        self.manifest[(m["item"], str(m.get("page", "")))] = m
            except (ValueError, TypeError):
                pass

    def match(self, text):
        flat = re.sub(r"\s+", "", text).lower()
        found = [(flat.find(i.lower()), i) for i in self.items if i.lower() in flat]
        return [i for _, i in sorted(found)]


# ---------------------------------------------------------------- field helpers

def holder_record(folder, pointer):
    if not pointer:
        return {}
    for p in sorted(folder.glob(f"sources/**/p{pointer}.json")):
        try:
            return json.loads(p.read_text(encoding="utf-8"))
        except ValueError:
            continue
    return {}


def record_title(rec):
    for k in ("title_dm", "title"):
        v = rec.get(k)
        if isinstance(v, str) and v.strip() and not v.strip().endswith("_"):
            return v.strip()
    return ""


def record_callid(rec):
    if isinstance(rec.get("callid"), str) and rec["callid"].strip():
        return rec["callid"].strip()
    f = rec.get("fields")
    if isinstance(f, str):
        m = re.search(r"'key': 'callid'.*?'value': '([^']+)'", f)
        if m:
            return m.group(1)
    elif isinstance(f, list):
        for x in f:
            if isinstance(x, dict) and x.get("key") == "callid":
                return str(x.get("value", ""))
    return ""


def pointers_from_header(header):
    fields = [f.strip() for f in header.split("|")]
    for f in fields[1:3]:
        if re.fullmatch(r"\d+(?:\s*[-,]\s*\d+)*", f):
            return re.findall(r"\d+", f)
    m = re.search(r"pointers? (\d+)(?:\s*[-,]\s*(\d+))?", header)
    return [g for g in m.groups() if g] if m else []


def page_from_header(header):
    m = re.search(r"\bPage (\d+(?:-\d+)?)", header) or re.search(r"\bp\.\s?(\d+(?:-\d+)?)", header)
    return ("Page " + m.group(1)) if m else ""


def call_from_text(text):
    m = re.search(r"\b(mss[A-Za-z]+ ?\d+[A-Za-z]*)", text)
    return m.group(1) if m else ""


def date_and_rest(text):
    m = DATE_RE.search(text)
    if not m:
        return "", text, ""
    after = re.sub(r"^\s*\([^)]*\)", "", text[m.end():])
    t = TIME_RE.match(after)
    when = m.group(1) + (", " + t.group(1).strip() if t else "")
    return when, text[:m.start()], after[t.end():] if t else after


def header_date(header):
    body = header.split("|")[-1]
    return date_and_rest(body)[0]


def header_place(header):
    body = re.sub(r"^\s*mss\w+ \d+ \([^)]*\),\s*", "", header.split("|")[-1].strip())
    _, _, after = date_and_rest(body)
    if not after or after.lstrip().startswith(","):
        return ""
    seg = after.split(",")[0].strip()
    if not seg or seg.lower().startswith(("to ", "for ", "operator")) or " to " in seg:
        return ""
    return seg


def ledger_place(lines):
    body = [l for l in lines if l.strip() and not l.startswith(NOTE_PREFIXES)][:3]
    for l in body:
        if WASH_RE.search(l):
            return "Washington"
        m = LEDGER_MONTH.search(l)
        if m:
            before = re.sub(r"\d{1,4}\s?(?:am|pm|a\.\s?m\.|p\.\s?m\.)", " ", l[:m.start()], flags=re.I)
            seg = re.split(r"\s{2,}", before.strip())[-1].strip(" ,.-") if before.strip() else ""
            if seg and len(seg.split()) <= 4 and not re.search(r"\d", seg):
                return seg
    return ""


def correspondents(cands):
    """(from, to): the first 'to' and the first non-empty 'from' over the candidates, in order. side 'before' reads
    the phrase before the date (a title or document text), 'after' the phrase after it (a block header), 'tail' the
    signature of a reading ('[signed] X' -> 'signed: X', used only when nothing names a sender)."""
    frm, to = "", ""
    for text, side in cands:
        if not text or (frm and to):
            continue
        if side == "tail":
            m = re.search(r"\{tail: \[signed\] ([^{}]*?)(?:\s+another\b|\s*\{|\}|$)", text)
            if m and not frm and m.group(1).strip():
                frm = "signature as written: " + m.group(1).strip()[:80]
            continue
        t = re.sub(r"^[^:]{0,60}:\s+", "", text) if side == "before" else text
        _, before, after = date_and_rest(t)
        seg = before if side == "before" else after
        seg = re.sub(r"\s*\((?:operator|signed|Cipher|LS|FM|row|page|image|transcription|sent|first|second)[^)]*\)",
                     "", seg).strip(" ,;")
        if not seg:
            continue
        if seg.lower().startswith("to "):
            f, t2 = "", seg[3:]
        elif " to " in seg:
            f, t2 = seg.split(" to ", 1)
        else:
            continue
        to = to or t2.strip(" ,")
        frm = frm or f.strip(" ,")
    return frm, to


def doc_parenthetical(doc):
    m = re.search(r"\(([^()]*(?:\([^()]*\)[^()]*)*)\)\s*$", doc)
    return m.group(1) if m else ""


def sentences(text):
    """Split on ';' and on '. X' at parenthesis depth 0, not after a known abbreviation."""
    out, cur, depth, i = [], [], 0, 0
    while i < len(text):
        ch = text[i]
        cur.append(ch)
        if ch == "(":
            depth += 1
        elif ch == ")":
            depth = max(0, depth - 1)
        elif depth == 0 and ch == ";":
            out.append("".join(cur[:-1]).strip())
            cur = []
        elif depth == 0 and ch == "." and i + 2 < len(text) and text[i + 1] == " " and text[i + 2].isupper():
            word = re.findall(r"[A-Za-z]+$", "".join(cur[:-1]))
            if not word or word[0].lower() not in ABBR:
                out.append("".join(cur).strip())
                cur = []
        i += 1
    if "".join(cur).strip():
        out.append("".join(cur).strip())
    return [s for s in out if s]


def split_top(s):
    """Split a sentence at ', its ' / ', the ' / ', and its ' outside parentheses only."""
    out, depth, start = [], 0, 0
    for i, ch in enumerate(s):
        depth += (ch == "(") - (ch == ")" and depth > 0)
        if depth == 0 and s.startswith(", ", i) and re.match(r",\s+(?:its |the |and its )", s[i:]):
            out.append(s[start:i])
            start = i + 2
    out.append(s[start:])
    return [x.strip() for x in out if x.strip()]


def prior_print_sentence(line):
    """One plain sentence: what was searched and that the item was not found there."""
    for s in sentences(line or ""):
        m = re.search(r"no prior [^;()]*?located\s*\((.*)\)\s*\.?$", s, flags=re.I)
        if m:
            what = m.group(1)
            d = re.search(r",?\s*(searched [^,]*)$", what)
            return "Not found in " + (what[:d.start()] + " (" + d.group(1) + ")" if d else what) + "."
        m = re.search(r"(?:no prior [^;]*?|\bnot )(?:was )?(?:located|found)\s+(?:(after [^;]*? of)|in|there, in)\s+(.*)",
                      s, flags=re.I)
        if m:
            what = m.group(2).rstrip(" .")
            if m.group(1):
                what = m.group(2).rstrip(" .") + " (" + m.group(1)[6:-3].strip() + ")"
            return "Not found in " + what + "."
        if re.search(r"\bprint neither\b|\bnot located\b|\bnot found\b", s, flags=re.I):
            return s.rstrip(".") + "."
    for s in sentences(line or ""):
        if re.search(r"\bno prior\b.*\b(located|found)\b", s, flags=re.I):
            return s[0].upper() + s[1:].rstrip(".") + "."
    return clean(line)


def print_clauses(rec, label="", others=()):
    """Positive 'in print' clauses from the verifier's fields (an outcome, reply, sequel or clear text printed); for a
    split row only the clauses about this item (or about no item of the row); near-duplicates dropped."""
    out = []
    for k in ("line", "grade", "gap", "unresolved_spans", "depth_check", "verifier_note", "plaintext_novelty_note"):
        v = rec.get(k)
        if not isinstance(v, str):
            continue
        for s in sentences(v):
            for part in split_top(s):
                part = re.sub(r"^[^:]{0,80}?\b(?:weak(?:est)?|N\d)\b[^:]{0,60}:\s*", "", part).strip()
                part = re.sub(r"^external(?: check)?(?: \([^)]*\))?:\s*", "", part).strip()
                part = re.sub(r"\s*\([^()]*\)", lambda m: "" if JARGON_RE.search(m.group(0)) else m.group(0), part)
                if (PRINT_RE.search(part) and not NEG_RE.search(part) and CITE_RE.search(part)
                        and not JARGON_RE.search(part) and not part.lower().startswith(("on ", "read at"))):
                    if not part or len(part) >= 400 or (others and not filter_unresolved(part, label, list(others))):
                        continue
                    part = part.rstrip(".") + "."
                    words = set(re.findall(r"[a-z0-9]{3,}", part.lower()))
                    dup = None
                    for j, o in enumerate(out):
                        ow = set(re.findall(r"[a-z0-9]{3,}", o.lower()))
                        if len(words & ow) >= 0.6 * max(1, min(len(words), len(ow))):
                            dup = j
                    if dup is None:
                        out.append(part)
                    elif len(part) > len(out[dup]):
                        out[dup] = part
    return " ".join(out[:3])


def weak_reason(rec, row, words, codes):
    if row.get("weak", "").lower() not in ("yes", "y", "true", "1"):
        return "no"
    text = " ".join(str(rec.get(k, "")) for k in ("gap", "line", "unresolved_spans", "verifier_note"))
    why = []
    if re.search(r"(clear|in clear)[^.;]{0,80}(transcription|since 2018)|public in clear|clear words are in", text, re.I):
        why.append("much of the message is already readable in clear in the holder's own volunteer transcription; "
                   "the key adds the code words")
    if re.search(r"\b(outcome|topic|exchange|order it answers|replies|reply|sequel|messages on the same)\b[^.;]{0,120}"
                 r"\b(printed|in print)\b|is in print", text, re.I):
        why.append("its outcome, reply or a related message is already in print (see 'partly in print?')")
    if not why:
        why.append("audit flag: the body is largely clear in the holder's transcription, or its topic is in print")
    if words and codes is not None:
        why.append(f"about {round(100 * (words - codes) / words)}% of its words are written in clear in the ledger")
    return "yes: " + "; ".join(why)


def class_cell(rec, row):
    g = rec.get("grade") or row.get("class") or ""
    m = re.search(r"\bN([0-5])\b", g)
    if not m:
        return clean(g)
    n = "N" + m.group(1)
    audits = clean(rec.get("audit_status", ""))
    k = re.match(r"(\w+) audits?", audits)
    tail = ""
    if k and NUMBER_WORDS.get(k.group(1).lower(), 0) >= 2:
        tail = f"; checked by {k.group(1).lower()} independent audits"
    return f"{n}: {CLASS_MEANING[n]}{tail}"


def depth_words(depth, pct):
    try:
        p = round(float(pct))
    except (TypeError, ValueError):
        p = None
    about = f" (about {p}%)" if p is not None else ""
    return {"D0": "not yet read", "D1": "fragments read", "D2": "partially deciphered" + about,
            "D3": "largely deciphered" + about, "D4": "deciphered"}.get(depth, "")


def grade_phrase(counts, unit="code words"):
    total = sum(counts.values())
    parts = [f"{counts[g]} {GRADE_LABEL[g]}" for g in GRADE_ORDER if counts.get(g)]
    return f"{total} {unit}: " + ", ".join(parts) if total else ""


def show(meaning):
    """A key value with alternatives ('s|que') shown as 's or que'."""
    return " or ".join(x for x in str(meaning).split("|")) if "|" in str(meaning) else str(meaning)


def pair_list(tokens):
    seen, order = {}, []
    for word, meaning, grade in tokens:
        k = (word.lower(), meaning, grade)
        if k not in seen:
            seen[k] = [word, meaning, grade, 0]
            order.append(k)
        seen[k][3] += 1
    out = []
    for k in order:
        word, meaning, grade, n = seen[k]
        if grade == "U" or meaning in ("?", ""):
            out.append(f"{word} (unread)" + (f" x{n}" if n > 1 else ""))
        else:
            out.append(f"{word} = {show(meaning)} ({GRADE_LABEL.get(grade, grade)})" + (f" x{n}" if n > 1 else ""))
    return "; ".join(out)


def dejargon(text):
    """Internal names in a reviewer's note, in words a holder can read."""
    text = re.sub(r"\s*and a G3 check", "", text)
    text = re.sub(r"\bkey(?:-no\d+)?\.md\b", "our key table", text)
    text = re.sub(r"\breading(?:-no\d+)?\.md\b", "our reading file", text)
    return text.replace("decoder H", "the decoder's cipher-book reading")


def plain_grades(text):
    """Spell out grade letters in a reviewer's note: '(M)' -> '(uncertain)', 'graded I' -> 'graded inferred'."""
    text = dejargon(text)
    lab = GRADE_LABEL
    text = re.sub(r"\(([HCSIMU])\)", lambda m: "(" + lab[m.group(1)] + ")", text)
    text = re.sub(r",\s*([HCSIMU])\)", lambda m: ", " + lab[m.group(1)] + ")", text)
    text = re.sub(r"\bgraded ([HCSIMU])\b", lambda m: "graded " + lab[m.group(1)], text)
    text = re.sub(r"\bone ([HCSIMU]) token", lambda m: "one token graded " + lab[m.group(1)], text)
    text = re.sub(r"\bword ([HCSIMU])\b(?!\.)", lambda m: "word graded " + lab[m.group(1)], text)
    text = re.sub(r"(?<=[\w'\]]) ([HCSIMU])(?=[\s;,)]|$)", lambda m: " (" + lab[m.group(1)] + ")", text)
    return text


def filter_unresolved(text, label, others):
    if not text or text.strip().lower() in ("none", "-"):
        return ""
    if not others:
        return text

    def hit(lab, clause):
        pat = r"\s*".join(re.escape(ch) for ch in lab.replace(" ", ""))
        return re.search(r"(?<![\w-])" + pat + r"(?![\w-]|\d)", clause, flags=re.I) is not None
    keep = [c.strip() for c in text.split(";") if c.strip()
            and (hit(label, c) or not any(hit(o, c) for o in others))]
    return "; ".join(keep)


# ---------------------------------------------------------------- selection drift

def selection_drift(status, listed_docs, classes, min_audits, scopes, folders):
    def passes(r):
        why = []
        m = re.search(r"\bN([0-5])\b", r.get("grade", ""))
        if classes and (not m or "N" + m.group(1) not in classes):
            why.append("class " + (("N" + m.group(1)) if m else "none"))
        k = re.match(r"(\w+) audits?", r.get("audit_status", "") or "")
        if min_audits and (not k or NUMBER_WORDS.get(k.group(1).lower(), 0) < min_audits):
            why.append("audits '" + (r.get("audit_status") or "") + "'")
        if scopes and r.get("claim_scope") not in scopes:
            why.append("scope " + str(r.get("claim_scope")))
        if r.get("superseded_by"):
            why.append("superseded")
        return why
    res = [r for r in status.get("results", []) if not folders or any(f in (r.get("link") or "") for f in folders)]
    by_doc = {r.get("document_id"): r for r in res}
    gone = []
    for d in listed_docs:
        r = by_doc.get(d)
        why = passes(r) if r else ["no status.json result"]
        if why:
            gone.append((d, ", ".join(why)))
    new = [r.get("document_id") for r in res if not passes(r) and r.get("document_id") not in listed_docs]
    return gone, new


# ---------------------------------------------------------------- build

def build_rows(list_rows, status, repo_root, meta, args):
    by_doc = {r.get("document_id"): r for r in status.get("results", [])}
    layouts, out = {}, []
    alias = dict(a.split("=", 1) for a in (args.alias or []))
    for row in list_rows:
        rec = by_doc.get(row["document"], {})
        if not rec:
            warn("no status.json result for: " + row["document"])
        folder = folder_path(row["folder"], repo_root)
        if folder not in layouts:
            lay = None
            try:
                if (folder / "decode.json").exists():
                    lay = DecodeKey(folder)
                elif list(folder.glob("reading*.md")):
                    lay = MdBlocks(folder)
            except Exception as e:  # a folder that cannot be read is warned, the row still written
                warn(f"cannot read {folder}: {e}")
            layouts[folder] = lay
        lay = layouts[folder]
        docs = rec.get("documents") or [row["document"]]
        items = []
        if lay:
            for d in docs:
                for i in lay.match(d):
                    if i not in items:
                        items.append(i)
        try:
            ndocs = int(row.get("documents") or len(docs))
        except ValueError:
            ndocs = len(docs)
        if not items:
            warn("no item found in the folder for: " + row["document"])
            items = [None]
        split = len(items) > 1
        if split and ndocs != len(items):
            warn(f"{row['document']}: documents {ndocs} but {len(items)} items found")
        for n, iid in enumerate(items):
            doc_for_item = next((d for d in docs if iid and lay and iid in lay.match(d)), row["document"])
            cells = item_cells(row, rec, lay, iid, doc_for_item, split, ndocs, items, alias, args)
            m = meta.get(iid or "", {})
            for k, i in META_KEYS.items():
                if m.get(k):
                    cells[i] = "" if m[k] == "-" else m[k]
            if m.get("note"):
                cells[9] = "Note: " + m["note"] + " -- " + cells[9]
            if not cells[3] and args.url_template and cells[2]:
                al = next((v for k, v in alias.items() if cells[0].startswith(k)), "")
                if al:
                    cells[3] = args.url_template.format(alias=al, pointer=cells[2])
            out.append([clean(c) for c in cells])
    return out


def item_cells(row, rec, lay, iid, doc, split, ndocs, items, alias, args):
    c = [""] * len(COLUMNS)
    title = rec.get("title", "")
    others = [i for i in items if i and i != iid]
    c[17] = row["folder"] + (f" (our entry {iid})" if iid else "")
    c[18] = args.credit
    c[16] = class_cell(rec, row)
    c[14] = dejargon(prior_print_sentence(rec.get("line", ""))) if rec else ""
    c[15] = print_clauses(rec, (iid or "").replace("BLA", "BLA "), [o.replace("BLA", "BLA ") for o in others]) \
        if rec else ""
    unresolved = filter_unresolved(plain_grades(row.get("unresolved", "")), (iid or "").replace("BLA", "BLA "), [
        o.replace("BLA", "BLA ") for o in others])
    depth = depth_words(row.get("depth (", ""), row.get("depth_pct", ""))
    if isinstance(lay, MdBlocks) and iid:
        it = lay.items[iid]
        ptrs = pointers_from_header(it["header"])
        recd = holder_record(lay.folder, ptrs[0] if ptrs else "")
        c[0] = record_callid(recd) or call_from_text(doc) or call_from_text(it["header"])
        page = record_title(recd) or page_from_header(it["header"])
        if len(ptrs) > 1:
            more = [record_title(holder_record(lay.folder, p)) or p for p in ptrs[1:]]
            page += "; continues on " + ", ".join(f"{t} (pointer {p})" for t, p in zip(more, ptrs[1:]))
        c[1] = page
        c[2] = ptrs[0] if ptrs else ""
        al = recd.get("collectionAlias") or next((v for k, v in alias.items() if c[0].startswith(k)), "")
        if args.url_template and c[2] and al:
            c[3] = args.url_template.format(alias=al, pointer=c[2])
        c[4] = header_date(it["header"]) or date_and_rest(title)[0]
        cands = ([(doc_parenthetical(doc), "before")] if split else [(title, "before")]) + [
            (s, "before") for s in summary_cells(lay.folder, iid)] + [(it["header"].split("|")[-1], "after"),
                                                                      (it["reading"], "tail")]
        c[5], c[6] = correspondents(cands)
        c[7] = header_place(it["header"]) or ledger_place(it["lines"])
        c[8] = lay.ledger_text(iid)
        toks = lay.tokens(iid)
        reading = it["reading"]
        if not split and ndocs > 1:
            reading = f"({ndocs} telegrams in this entry) " + reading
        c[9] = reading
        c[10] = pair_list(toks)
        mine = {}
        for _, _, g in toks:
            mine[g] = mine.get(g, 0) + 1
        if it["counts"] and mine != it["counts"]:
            warn(f"{iid}: token pass {mine} differs from the reading's count line {it['counts']}")
        unc = [f"{w} = {m} ({GRADE_LABEL.get(g, g)})" for w, m, g in toks if g in ("M", "I", "U")]
        c[11] = "; ".join(dict.fromkeys(unc + ([unresolved] if unresolved else []))) or "none"
        counts = it["counts"] or mine
        c[12] = "; ".join(x for x in (depth + f" ({row.get('depth (', '')})" if depth else "",
                                      grade_phrase(counts)) if x)
        c[13] = weak_reason(rec, row, lay.word_count(iid), sum(counts.values()))
    elif isinstance(lay, DecodeKey) and iid:
        it = lay.items[iid]
        c[0] = call_from_text(doc)
        pages = [(p, lay.manifest.get((iid, p), {}).get("pointer", "")) for p in it["pages"]]
        c[1] = ", ".join(f"{p} (page pointer {ptr})" if ptr else p for p, ptr in pages)
        c[4] = date_and_rest(doc)[0]
        c[8] = " / ".join(" ".join(g) for _, g in sorted(it["cipher"].items()))
        c[9] = " / ".join(it["reading"])
        c[10] = pair_list([(g, v, gr) for g, v, gr in it["tokens"]])
        unc = [f"{g} (unread)" if gr == "U" else f"{g} = {show(v)} ({GRADE_LABEL.get(gr, gr)})"
               for g, v, gr in it["tokens"] if gr in ("M", "I", "U")]
        c[11] = "; ".join(dict.fromkeys(unc + ([unresolved] if unresolved else []))) or "none"
        counts = {}
        for _, _, gr in it["tokens"]:
            counts[gr] = counts.get(gr, 0) + 1
        whole = depth + f" ({row.get('depth (', '')})" if depth else ""
        c[12] = "; ".join(x for x in (grade_phrase(counts, "cipher groups"),
                                      (f"for all {len(items)} items of this entry together: " if split else "")
                                      + whole) if x)
        c[13] = weak_reason(rec, row, 0, None)
    else:
        c[0] = call_from_text(row["document"])
        c[9] = row.get("title", "")
        c[12] = "; ".join(x for x in (depth, row.get("completeness", "")) if x)
        c[13] = weak_reason(rec, row, 0, None)
    if not c[5] and not c[6] and not isinstance(lay, MdBlocks):
        c[5], c[6] = correspondents([(doc_parenthetical(doc), "before"), (title, "before")])
    return c


_SUMMARY_CACHE = {}


def summary_cells(folder, iid):
    """Second column of a '| ID | sender, addressee, date | ...' summary table row in the folder's reading*.md."""
    if folder not in _SUMMARY_CACHE:
        d = {}
        for p in sorted(folder.glob("reading*.md")):
            for l in p.read_text(encoding="utf-8").splitlines():
                m = re.match(r"^\|\s*([A-Za-z0-9-]+)\s*\|\s*([^|]*)\|", l)
                if m:
                    d.setdefault(m.group(1), []).append(m.group(2).strip())
        _SUMMARY_CACHE[folder] = d
    return _SUMMARY_CACHE[folder].get(iid, [])


# ---------------------------------------------------------------- output

def readme_rows(holder, credit):
    h = holder or "the holder"
    return [
        ("What this file is", f"One row per item (telegram, letter or enclosure) that Cipher Lab has read from {h}'s "
         "collections, keyed by the holder's own page or object number so that each row can be matched to an "
         "existing catalogue record. Every reading is regenerated by a script from our transcription and the key, "
         "and every code word is graded. Nothing here claims priority: 'not found in' means only that the "
         "editions named were searched and the text was not found there."),
        ("What is ours and what is theirs", f"The page images, the catalogue records and the volunteer transcription "
         f"on {h}'s digital library are {h}'s. The transcriptions in this file were made from those page images, "
         "with the volunteer transcription used as a further witness (a few entries, or parts of them, were taken from "
         "the volunteer transcription alone; the repository's notes on each entry say so). The readings, the code-word lists, the keys "
         "and the search notes are Cipher Lab's and are public under the repository's MIT licence "
         f"({REPO_URL}); a credit in the record would be welcome but is not required. Suggested: '{credit}'."),
        ("How to import", "The .tsv file is tab-delimited UTF-8 text with one header row and no tab or line break "
         "inside any cell; the .csv holds the same rows. The column 'holder pointer (CONTENTdm number)' is the "
         "holder's own number for the page (or for the whole item where a page list is given), so rows can be "
         "matched to records by that number. Map only the columns you want (for example reading, code words, "
         "credit) to your own fields; the others are there for checking. Where an entry runs onto a second page "
         "the row is keyed to the first page and the 'holder page' cell names the continuation."),
        ("Grades of a code word", "from the cipher book = the meaning is written in the period cipher book or key; "
         "from a printed or period copy = only a printed edition or a period decipherment gives it; cryptanalytic = "
         "worked out from the cipher itself with a control test; inferred = fixed from another message; uncertain = "
         "a reading we are not sure of; unread = a code word or group we could not read."),
        ("Classes (our class)", "; ".join(f"{k} = {v}" for k, v in CLASS_MEANING.items()) + ". Each class was set "
         "by a reviewing session separate from the one that made the reading."),
        ("How much is read", "deciphered = every cipher token read and checked against an outside source; largely "
         "deciphered (about N%) = at least 80% of the tokens read with a check; partially deciphered (about N%) = "
         "at least one longer passage reads and one true sentence about the content can be written, but the "
         "checks for 'largely' are not all met, so an item can have every code word read and still be called "
         "partially deciphered; fragments read = scattered words only. The D-number is our depth scale (D1-D4)."),
        ("Adds less? (and why)", "yes = a reviewer flagged that our reading adds less than usual: much of the "
         "message is already in clear in the holder's own transcription (the clerks wrote many plain words), or "
         "its outcome, reply or a related message is already in print. The share of words written in clear is "
         "counted from our transcription."),
    ] + [(col, COLUMN_HELP[col]) for col in COLUMNS]


COLUMN_HELP = {
    "holder call number": "The holder's call number for the volume or item.",
    "holder page (as the holder titles it)": "The page title in the holder's digital library record; for a "
    "multi-page item, our page labels with each page's own number.",
    "holder pointer (CONTENTdm number)": "The holder's record number for the page (or item); the key to import by.",
    "holder URL": "The record's page in the holder's digital library.",
    "date": "The date (and time) the ledger or letter gives.",
    "from": "The sender, as the reviewers name them (an operator or forwarding office may be named with 'via').",
    "to": "The addressee, as the reviewers name them.",
    "place": "The place the ledger heads the entry with ('Washn', 'Wash'n' and the like written out as Washington); "
    "blank where the page gives none.",
    "ledger/page text as we transcribed it": "Our transcription, lines separated by ' / '. <del>..</del> = struck "
    "through, <ins>..</ins> = written above the line, word[?] = doubtful, ' = ' joins a word the clerk split. For "
    "number ciphers, the cipher groups line by line.",
    "reading": "The decoded text: [..] = a code word replaced by its meaning, {time: ..} and {date: ..} = time and "
    "date code words, {tail: ..} = everything from the signature code word on (signature, service notes and, where "
    "the ledger has one, a second telegram). For number ciphers, [n] = a group not in the key.",
    "code words and meanings": "Each code word or cipher group as written = its meaning (grade); 'xN' = it occurs N "
    "times. Punctuation and signature code words are included.",
    "uncertain or unread words": "Code words graded uncertain or inferred, and anything left unread, with the "
    "reviewers' own note.",
    "how much is read": "Depth in words (see above) and the count of code words by grade.",
    "adds less? (and why)": "See above.",
    "prior print checked": "Where the reviewers searched; 'not found in' is a search result, not a claim of priority.",
    "partly in print?": "Related text that is in print (an outcome, a reply, a sequel, or the clear part of the "
    "letter), as the reviewers name it; blank if they name none.",
    "our class": "See 'Classes' above.",
    "repository link": "The folder with the transcription, key, reading and search log (AUDIT.md); our entry id in "
    "brackets.",
    "suggested credit line": "A credit line, if you wish to use one.",
}


def write_text(rows, stem):
    """STEM.csv (quoted where needed) and STEM.tsv (plain tab-joined; clean() keeps tabs and breaks out of cells)."""
    with open(f"{stem}.csv", "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(COLUMNS)
        w.writerows(rows)
    with open(f"{stem}.tsv", "w", encoding="utf-8", newline="") as f:
        for r in [COLUMNS] + rows:
            f.write("\t".join(clean(c) for c in r) + "\n")


def write_xlsx(rows, stem, holder, credit):
    """Return True when the .xlsx was written, False when openpyxl is not importable."""
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font
        from openpyxl.utils import get_column_letter
    except ImportError:
        return False
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Readings"
    ws.append(COLUMNS)
    for r in rows:
        ws.append(r)
    widths = {0: 12, 1: 22, 2: 12, 3: 30, 4: 16, 5: 22, 6: 22, 7: 14, 8: 60, 9: 70, 10: 60, 11: 40, 12: 34, 13: 40,
              14: 50, 15: 40, 16: 40, 17: 40, 18: 30}
    for i in range(len(COLUMNS)):
        ws.column_dimensions[get_column_letter(i + 1)].width = widths.get(i, 20)
    wrap = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = wrap
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for r in range(2, ws.max_row + 1):
        for col in (4, 18):
            cell = ws.cell(row=r, column=col)
            url = str(cell.value or "").split(" ")[0]
            if url.startswith("http"):
                cell.hyperlink = url
    rd = wb.create_sheet("README")
    rd.append(["Topic", "Explanation"])
    for k, v in readme_rows(holder, credit):
        rd.append([k, v])
    rd.column_dimensions["A"].width = 34
    rd.column_dimensions["B"].width = 120
    for row in rd.iter_rows():
        for cell in row:
            cell.alignment = wrap
    for cell in rd[1]:
        cell.font = Font(bold=True)
    rd.freeze_panes = "A2"
    wb.save(f"{stem}.xlsx")
    return True


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0], formatter_class=argparse.RawDescriptionHelpFormatter,
                                 epilog=__doc__.split("\n\n", 1)[1])
    ap.add_argument("list", help="list TSV (one row per entry; see the docstring for the columns)")
    ap.add_argument("--out", required=True, help="output path without extension (writes .csv, .tsv and .xlsx)")
    ap.add_argument("--status", default=str(ROOT / "status.json"))
    ap.add_argument("--repo-root", default=str(ROOT))
    ap.add_argument("--meta", help="per-item overrides TSV (item, holder_call, holder_page, pointer, url, date, from, "
                    "to, place, note)")
    ap.add_argument("--url-template", default="", help="e.g. https://host/digital/collection/{alias}/id/{pointer}")
    ap.add_argument("--alias", action="append", help="CALLPREFIX=collection alias, e.g. mssEC=p16003coll11")
    ap.add_argument("--holder", default="", help="the holder's name for the README sheet")
    ap.add_argument("--credit", default=f"Decipherment: Cipher Lab, "
                    f"{datetime.datetime.now(datetime.timezone.utc).year}, {REPO_URL}")
    ap.add_argument("--no-xlsx", action="store_true")
    ap.add_argument("--select-class", default="", help="comma list, e.g. N3,N4")
    ap.add_argument("--select-min-audits", type=int, default=0)
    ap.add_argument("--select-scope", default="")
    ap.add_argument("--select-folder", default="", help="comma list of folder names a result's link must contain")
    args = ap.parse_args(argv)

    list_rows = read_list(args.list)
    status = json.loads(Path(args.status).read_text(encoding="utf-8"))
    meta = read_meta(args.meta)
    if args.select_class or args.select_min_audits or args.select_scope or args.select_folder:
        split = (lambda s: [x.strip() for x in s.split(",") if x.strip()])
        gone, new = selection_drift(status, [r["document"] for r in list_rows], split(args.select_class),
                                    args.select_min_audits, split(args.select_scope), split(args.select_folder))
        print(f"selection drift: {len(gone)} listed row(s) no longer pass, {len(new)} passing result(s) not listed")
        for d, why in gone:
            print(f"  listed, now fails ({why}): {d}")
        for d in new:
            print(f"  passes, not listed: {d}")
    rows = build_rows(list_rows, status, args.repo_root, meta, args)
    Path(args.out).parent.mkdir(parents=True, exist_ok=True)
    write_text(rows, args.out)
    made = []
    if not args.no_xlsx:
        if write_xlsx(rows, args.out, args.holder, args.credit):
            made.append("xlsx")
        else:
            warn("openpyxl not importable: wrote .csv and .tsv only (pip install openpyxl for the .xlsx)")
    print(f"wrote {len(rows)} rows to {args.out}.csv, .tsv" + (", .xlsx" if made else ""))
    return 0


if __name__ == "__main__":
    sys.exit(main())
