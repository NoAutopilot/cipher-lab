"""Offline tests for tools/holder_export.py (HOLDER-EXPORT, 9 Oct 2026).

Must catch: a change in the column order a holder imports by; a list row naming two telegrams the reading keeps apart
not split into two rows; a ledger entry holding two telegrams not labelled '2 telegrams in this entry'; the .xlsx
step failing the whole run when openpyxl is missing; '[SIGN-OFF]' or an e-mail address (rule 9) leaking from
status.json into any output; a grade letter left unexplained; a tab or line break inside a cell.
Must NOT block: a list row with no status.json result (written from the list's own fields, warned only).
Fixture: tools/tests/fixtures/holder_export/ (three list rows, a toy ledger with its own decode.py, a number cipher in
the tools/decode_key.py layout).
"""
import csv
import io
import json
import re
import shutil
import sys
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))
import holder_export as he  # noqa: E402

FIX = HERE / "fixtures" / "holder_export"
EXPECTED = [
    "holder call number", "holder page (as the holder titles it)", "holder pointer (CONTENTdm number)", "holder URL",
    "date", "from", "to", "place", "ledger/page text as we transcribed it", "reading", "code words and meanings",
    "uncertain or unread words", "how much is read", "adds less? (and why)", "prior print checked", "partly in print?",
    "our class", "repository link", "suggested credit line",
]
EMAIL = re.compile(r"[\w.+-]+@[\w-]+\.[\w.]+")


def run(out, status=None, extra=()):
    argv = [str(FIX / "list.tsv"), "--out", str(out), "--status", str(status or FIX / "status.json"),
            "--repo-root", str(FIX), "--url-template", "https://example.org/{alias}/id/{pointer}",
            "--alias", "Demo=demo1", "--alias", "mssDN=demo2", "--holder", "the Demo Library",
            "--select-class", "N3,N4", "--select-min-audits", "2"] + list(extra)
    so, se = io.StringIO(), io.StringIO()
    with redirect_stdout(so), redirect_stderr(se):
        rc = he.main(argv)
    return rc, so.getvalue(), se.getvalue()


def rows_of(stem):
    with open(f"{stem}.csv", encoding="utf-8") as f:
        rows = list(csv.reader(f))
    return rows[0], [dict(zip(rows[0], r)) for r in rows[1:]]


class TestHolderExport(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.stem = self.tmp / "out" / "demo"

    def tearDown(self):
        shutil.rmtree(self.tmp)

    def test_column_order(self):
        self.assertEqual(he.COLUMNS, EXPECTED)
        run(self.stem, extra=["--no-xlsx"])
        head, _ = rows_of(self.stem)
        self.assertEqual(head, EXPECTED)
        tsv = Path(f"{self.stem}.tsv").read_text(encoding="utf-8").splitlines()
        self.assertEqual(tsv[0].split("\t"), EXPECTED)
        for line in tsv:  # no tab or line break inside a cell: every line has exactly the header's field count
            self.assertEqual(len(line.split("\t")), len(EXPECTED))

    def test_split_and_two_telegram_note(self):
        rc, out, _ = run(self.stem, extra=["--no-xlsx"])
        self.assertEqual(rc, 0)
        _, rows = rows_of(self.stem)
        self.assertEqual(len(rows), 4)  # T1 and T2 split, T3 one row, DN7 one row
        ids = [r["repository link"].split("our entry ")[-1].rstrip(")") for r in rows]
        self.assertEqual(ids, ["T1", "T2", "T3", "DN7"])
        t1, t2, t3, dn = rows
        self.assertEqual((t1["from"], t1["to"]), ("Alpha", "Beta"))
        self.assertEqual((t2["from"], t2["to"]), ("Gamma", "Beta"))
        self.assertEqual(t1["date"], "1 Jan 1864, 9 AM")
        self.assertEqual(t2["date"], "2 Jan 1864, 10 AM")
        self.assertFalse(t1["reading"].startswith("("))
        self.assertTrue(t3["reading"].startswith("(2 telegrams in this entry)"), t3["reading"])
        # holder keys: page title from the cached record, first pointer, continuation named
        self.assertEqual(t1["holder pointer (CONTENTdm number)"], "101")
        self.assertEqual(t1["holder call number"], "Demo vol 1")
        self.assertEqual(t1["holder URL"], "https://example.org/demo1/id/101")
        self.assertEqual(t3["holder page (as the holder titles it)"], "Page 2; continues on Page 3 (pointer 103)")
        self.assertEqual(t1["place"], "Washington")
        # ledger lines kept apart with ' / ', orchestrator notes dropped
        self.assertIn(" / ", t1["ledger/page text as we transcribed it"])
        self.assertNotIn("note:", t1["ledger/page text as we transcribed it"])
        # decode-key layout: page pointer from the manifest, groups line by line
        self.assertEqual(dn["holder page (as the holder titles it)"], "p1 (page pointer 701)")
        self.assertEqual(dn["ledger/page text as we transcribed it"], "10 11 / 12 13")
        self.assertIn("13 = part or parti (uncertain)", dn["code words and meanings"])
        self.assertIn("12 (unread)", dn["uncertain or unread words"])
        # selection drift is reported (DN7 has one audit), not fatal
        self.assertIn("listed, now fails", out)

    def test_grades_and_wording(self):
        run(self.stem, extra=["--no-xlsx"])
        _, rows = rows_of(self.stem)
        t1 = rows[0]
        self.assertIn("Apple = General Beta (from the cipher book)", t1["code words and meanings"])
        self.assertIn("Plum = Richmond (uncertain)", t1["uncertain or unread words"])
        self.assertIn("Lime = Alpha (inferred)", t1["uncertain or unread words"])
        self.assertEqual(t1["how much is read"],
                         "deciphered (D4); 5 code words: 3 from the cipher book, 1 inferred, 1 uncertain")
        self.assertTrue(t1["adds less? (and why)"].startswith("yes: much of the message"))
        self.assertTrue(t1["prior print checked"].startswith("Not found in the Demo Records"))
        self.assertTrue(t1["our class"].startswith("N4: no prior decipherment located"))
        self.assertEqual(rows[2]["partly in print?"], "its reply is printed in Demo Records vol. 2 p.7.")
        for r in rows:
            for col in ("code words and meanings", "uncertain or unread words", "how much is read"):
                self.assertIsNone(re.search(r"\([HCSIMU]\)", r[col]), (col, r[col]))
            # rule 10: no priority claim (the N4 meaning's own caveat 'unpublished work not excluded' is allowed)
            text = " ".join(r.values()).lower()
            self.assertIsNone(re.search(r"first decipherment|previously unread|newly recovered|never printed|"
                                        r"unpublished plaintext|\bis unpublished", text), text)

    def test_no_signoff_or_email_leak(self):
        run(self.stem)
        for ext in ("csv", "tsv"):
            text = Path(f"{self.stem}.{ext}").read_text(encoding="utf-8")
            self.assertNotIn("[SIGN-OFF]", text)
            self.assertIsNone(EMAIL.search(text), ext)
        if Path(f"{self.stem}.xlsx").exists():
            import openpyxl
            wb = openpyxl.load_workbook(f"{self.stem}.xlsx")
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    for v in row:
                        self.assertNotIn("[SIGN-OFF]", str(v or ""))
                        self.assertIsNone(EMAIL.search(str(v or "")), v)

    def test_fallback_without_openpyxl(self):
        saved = {k: sys.modules[k] for k in list(sys.modules) if k == "openpyxl" or k.startswith("openpyxl.")}
        sys.modules["openpyxl"] = None  # import openpyxl now raises ImportError
        try:
            rc, out, err = run(self.stem)
        finally:
            del sys.modules["openpyxl"]
            sys.modules.update(saved)
        self.assertEqual(rc, 0)
        self.assertTrue(Path(f"{self.stem}.csv").exists())
        self.assertTrue(Path(f"{self.stem}.tsv").exists())
        self.assertFalse(Path(f"{self.stem}.xlsx").exists())
        self.assertIn("openpyxl not importable", err)

    def test_xlsx_when_available(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl not installed")
        run(self.stem)
        wb = openpyxl.load_workbook(f"{self.stem}.xlsx")
        self.assertEqual(wb.sheetnames, ["Readings", "README"])
        ws = wb["Readings"]
        self.assertEqual([c.value for c in ws[1]], EXPECTED)
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertTrue(ws.cell(row=2, column=10).alignment.wrap_text)
        topics = [r[0] for r in wb["README"].iter_rows(min_row=2, values_only=True)]
        for col in EXPECTED:  # every column is explained
            self.assertIn(col, topics)
        self.assertIn("How to import", topics)
        self.assertIn("What is ours and what is theirs", topics)

    def test_row_without_status_is_written(self):
        st = json.loads((FIX / "status.json").read_text(encoding="utf-8"))
        st["results"] = st["results"][1:]  # drop T1/T2's result
        sp = self.tmp / "status.json"
        sp.write_text(json.dumps(st), encoding="utf-8")
        rc, _, err = run(self.stem, status=sp, extra=["--no-xlsx"])
        self.assertEqual(rc, 0)
        self.assertIn("no status.json result", err)
        _, rows = rows_of(self.stem)
        self.assertEqual(rows[0]["repository link"].split("our entry ")[-1].rstrip(")"), "T1")

    def test_meta_override(self):
        meta = self.tmp / "meta.tsv"
        meta.write_text("item\tpointer\tfrom\tin_print\tnote\tsource\nDN7\t700\tsigned 'N'\t-\tlines 1-2 only\tfixture\n",
                        encoding="utf-8")
        run(self.stem, extra=["--no-xlsx", "--meta", str(meta)])
        _, rows = rows_of(self.stem)
        dn = rows[-1]
        self.assertEqual(dn["holder pointer (CONTENTdm number)"], "700")
        self.assertEqual(dn["holder URL"], "https://example.org/demo2/id/700")
        self.assertEqual(dn["from"], "signed 'N'")
        self.assertEqual(dn["partly in print?"], "")
        self.assertTrue(dn["reading"].startswith("Note: lines 1-2 only -- "))

    def test_sentences_keep_abbreviations(self):
        s = he.sentences("not located in OR ser. I vol. 34 pt 4 (searched 7 Oct 2026); its reply is printed. Next one")
        self.assertEqual(s[0], "not located in OR ser. I vol. 34 pt 4 (searched 7 Oct 2026)")
        self.assertEqual(len(s), 3)


if __name__ == "__main__":
    unittest.main()
